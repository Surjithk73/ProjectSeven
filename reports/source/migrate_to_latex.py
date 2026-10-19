"""One-time migration. Do not rerun over subsequent edits to the LaTeX master."""
from pathlib import Path
import json
import re
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'source'
OUT = ROOT / 'APIShield_Project_Report.tex'
doc = Document(ROOT / 'APIShield_Project_Report.docx')

def esc(s):
    s = s.replace('\ufffd', '"')
    substitutions = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '$': r'\$', '#': r'\#', '_': r'\_', '{': r'\{', '}': r'\}', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}', '→': r'\(\rightarrow\)', '±': r'\(\pm\)', 'κ': r'\(\kappa\)', '≤': r'\(\leq\)', '≥': r'\(\geq\)'}
    return ''.join(substitutions.get(c, c) for c in s)

preamble = r'''% APIShield maintained report master. Edit this file directly from now on.
% Previous DOCX and Markdown files are historical exports, not input sources.
% Compile with XeLaTeX (preferred) or pdfLaTeX; compile twice for page references.
% All seven system diagrams are inline TikZ; references are included below.
\documentclass[12pt,a4paper,oneside]{report}
\usepackage[left=3.81cm,right=2.54cm,top=2.54cm,bottom=2.54cm]{geometry}
\usepackage{iftex}
\ifPDFTeX
  \usepackage[T1]{fontenc}
  \usepackage[utf8]{inputenc}
  \usepackage{newtxtext,newtxmath}
\else
  \usepackage{fontspec}
  \IfFontExistsTF{Times New Roman}{\setmainfont{Times New Roman}}{\setmainfont{TeX Gyre Termes}}
\fi
\usepackage{setspace,titlesec,longtable,array,booktabs,graphicx,caption}
\usepackage{tikz}
\usetikzlibrary{arrows.meta,positioning}
\usepackage{fancyvrb,fancyhdr,xurl}
\usepackage[hidelinks,unicode]{hyperref}
\hypersetup{pdftitle={APIShield Intent-Aware Detection and Validated Repair of Breaking API Changes},pdfauthor={[CANDIDATE NAME]}}
\onehalfspacing
\setlength{\parindent}{0pt}
\setlength{\parskip}{4pt}
\setlength{\emergencystretch}{3em}
\widowpenalty=10000
\clubpenalty=10000
\setcounter{tocdepth}{1}
\setcounter{secnumdepth}{1}
\titleformat{\chapter}[display]{\centering\bfseries\fontsize{16}{20}\selectfont}{\MakeUppercase{\chaptertitlename}\ \thechapter}{8pt}{\MakeUppercase}
\titlespacing*{\chapter}{0pt}{0pt}{18pt}
\titleformat{\section}{\normalfont\fontsize{14}{17}\selectfont}{\thesection}{0.6em}{\MakeUppercase}
\titlespacing*{\section}{0pt}{12pt}{8pt}
\captionsetup{font=small,labelfont=normalfont,labelsep=space}
\pagestyle{fancy}
\fancyhf{}\fancyfoot[C]{\thepage}
\renewcommand{\headrulewidth}{0pt}
\fancypagestyle{plain}{\fancyhf{}\fancyfoot[C]{\thepage}\renewcommand{\headrulewidth}{0pt}}
\renewcommand{\contentsname}{TABLE OF CONTENTS}
\renewcommand{\listfigurename}{LIST OF FIGURES}
\renewcommand{\listtablename}{LIST OF TABLES}
\newcommand{\frontchapter}[1]{\chapter*{#1}\addcontentsline{toc}{chapter}{#1}}
\newcommand{\screenshotplaceholder}[2]{% optional: insert only when a figure is planned
  \begin{figure}[htbp]\centering
  \fbox{\parbox[c][45mm][c]{0.9\linewidth}{\centering #1}}
  \caption{#2}\end{figure}}
\tikzset{reportbox/.style={draw,rounded corners=2pt,fill=black!3,align=center,font=\small,minimum height=13mm,text width=27mm},reportarrow/.style={-{Stealth[length=2mm]},semithick}}
\begin{document}
\pagenumbering{arabic}
'''
out = [preamble]

# Cover/title and front matter are preserved from the delivered report.
for start, stop in [(0,10),(10,20)]:
    out.append(r'\begin{center}')
    for p in doc.paragraphs[start:stop]:
        t=p.text.strip()
        if not t: continue
        if p.style.name=='Title': out.append(r'{\bfseries\fontsize{18}{23}\selectfont '+esc(t)+r'\par}\vspace{10mm}')
        else: out.append(esc(t).replace('\n',r'\\'+'\n')+r'\par\vspace{4mm}')
        if t=='by': pass
        if 'TEAM MEMBER 3' in t:
            out.append(r'\IfFileExists{assets/institution_logo.jpg}{\includegraphics[width=23mm]{assets/institution_logo.jpg}\par\vspace{3mm}}{}')
    out.append(r'\end{center}\clearpage')
for p in doc.paragraphs[20:42]:
    if not p.text.strip(): continue
    if p.style.name=='Front Heading': out.append(r'\frontchapter{'+esc(p.text)+'}')
    else: out.append(esc(p.text).replace('\n',r'\\'+'\n')+'\par\n')
out.extend([r'\clearpage\tableofcontents',r'\clearpage\listoffigures',r'\clearpage\listoftables',r'\frontchapter{LIST OF SYMBOLS AND ABBREVIATIONS}'])
active=False
for p in doc.paragraphs:
    if p.text=='LIST OF SYMBOLS AND ABBREVIATIONS': active=True; continue
    if active and p.text.startswith('CHAPTER 1'): break
    if active and p.text.strip(): out.append(esc(p.text.replace('\t',' — '))+r'\par')

def tikz(kind):
    if kind in ('workflow','publication','pipeline'):
        labels={
            'workflow':['Connect project\\and contracts','Check changes\\and impact','Confirm rules\\and checks','Request repair\\and review','Patch or draft\\preview'],
            'publication':['Validated\\retained patch','Offline\\revalidation','Revisioned\\full preview','Explicit\\draft publish'],
            'pipeline':['Captured source\\and contracts','Static scan\\and comparison','Guarded\\candidate edits','Configured tests\\and rule checks']
        }[kind]
        cmds=[]
        for n,label in enumerate(labels):
            cmds.append(r'\node[reportbox] (n'+str(n)+') at ('+str(n*3.25)+',0) {'+label+'};')
            if n: cmds.append(r'\draw[reportarrow] (n'+str(n-1)+') -- (n'+str(n)+');')
        if kind=='publication': cmds.append(r'\node[align=center,font=\small] at (4.875,-1.5) {Stale state or lost response: reconcile before publication};')
        if kind=='pipeline': cmds.append(r'\node[align=center,font=\small] at (4.875,-1.5) {Each retry starts from the same captured baseline};')
        if kind=='workflow': cmds.append(r'\node[align=center,font=\small] at (6.5,-1.5) {Generation and publication require separate explicit actions};')
        return '\n'.join(cmds)
    specs={
      'architecture':([('browser',0,2,'React browser'),('cli',0,0,'CLI / research\\harness'),('api',4,2,'FastAPI\\coordination'),('engine',4,0,'Shared Python\\engine'),('records',8,2,'Durable local\\evidence'),('external',8,0,'oasdiff / Gemini\\GitHub')],[('browser','api'),('cli','engine'),('api','engine'),('api','records'),('engine','external'),('engine','records')]),
      'states':([('queued',0,0,'Queued'),('running',3.5,0,'Running'),('completed',8,2.7,'Completed'),('validated',8,.9,'Validated'),('input',8,-.9,'Needs input'),('failed',8,-2.7,'Failed / cancelled')],[('queued','running'),('running','completed'),('running','validated'),('running','input'),('running','failed')]),
      'graph':([('caller',0,2,'Caller function'),('consumer',4,2,'HTTP consumer'),('operation',8,2,'API operation'),('rule',4,0,'Confirmed rule'),('test',8,0,'Committed test')],[('caller','consumer'),('consumer','operation'),('consumer','rule'),('rule','test')]),
      'oracle':([('context',0,2,'Contracts and source\\Rules / public feedback'),('generation',5,2,'Candidate generation\\Engine acceptance'),('verify',5,-.3,'Separate verification\\Hidden behavioral checks')],[('context','generation'),('generation','verify')])
    }
    nodes,edges=specs[kind]
    s=[r'\node[reportbox] ('+name+') at ('+str(x)+','+str(y)+') {'+text+'};' for name,x,y,text in nodes]
    for a,b in edges:
        dashed=',dashed' if kind=='graph' and (b=='rule' or a=='rule') else ''
        s.append(r'\draw[reportarrow'+dashed+'] ('+a+') -- ('+b+');')
    if kind=='graph': s.append(r'\node[align=center,font=\small] at (4,-1.5) {Solid: static observation. Dashed: developer declaration.};')
    if kind=='oracle': s.extend([r'\draw[dashed] (-1.6,.9) -- (6.7,.9);',r'\node[align=center,font=\small,text width=35mm] at (0,-.3) {Hidden checks never enter prompts or retry feedback};'])
    if kind=='states': s.append(r'\node[align=center,font=\small] at (3.5,-3.6) {No change is a separate repair outcome};')
    return '\n'.join(s)

lines=(SRC/'report.md').read_text(encoding='utf-8').splitlines()
refs=json.loads((SRC/'references.json').read_text(encoding='utf-8'))
i=0; refs_added=False; figures=0; tables=0
while i<len(lines):
    line=lines[i]
    if line=='<!-- PAGE -->': i+=1; continue
    if line.startswith('# CHAPTER ') or line.startswith('# APPENDIX '):
        appendix=line.startswith('# APPENDIX ')
        n=int(line.rsplit(' ',1)[1]); title=lines[i+1][2:]
        if appendix:
            out.extend([r'\renewcommand{\chaptername}{APPENDIX}',r'\setcounter{chapter}{'+str(n-1)+'}'])
        out.append(r'\chapter{'+esc(title)+'}')
        i+=2; continue
    if line.startswith('## '):
        title=re.sub(r'^\d+\.\d+\s+','',line[3:])
        out.append(r'\section{'+esc(title)+'}'); i+=1; continue
    if line=='# REFERENCES':
        out.append(r'\frontchapter{REFERENCES}')
        for n,ref in enumerate(refs,1):
            url=ref['url']; link=r'\url{'+url+'}' if url.startswith('http') else esc('ProjectSeven repository, '+url+'.')
            out.append(r'\begingroup\setlength{\hangindent}{1.27cm}\setlength{\hangafter}{1} '+esc(ref['author'].rstrip('.')+'. ('+ref['year']+'). ')+r'\emph{'+esc(ref['title'])+'}. '+esc(ref['type'])+'. '+link+r'\par\endgroup\vspace{6pt}')
        refs_added=True; i+=1; continue
    if line.startswith('[[REFERENCES:'): i+=1; continue
    if line.startswith('[[TABLE:'):
        number,title,*cols=line[8:-2].split('|'); rows=[]; i+=1
        while lines[i]!='[[ENDTABLE]]': rows.append(lines[i].split('|')); i+=1
        fractions={2:[.34,.66],3:[.27,.38,.35],4:[.27,.24,.27,.22]}[len(cols)]
        if number=='A6.1': fractions=[.52,.48]
        spec='@{}'+''.join(r'>{\raggedright\arraybackslash}p{'+f'{v:.3f}'+r'\dimexpr\linewidth-'+str(2*(len(cols)-1))+r'\tabcolsep\relax}' for v in fractions)+'@{}'
        # Subtract intercolumn space once from the total usable width.
        spec='@{}'+''.join(r'>{\raggedright\arraybackslash}p{\dimexpr '+f'{v:.3f}'+r'\linewidth-'+f'{v*2*(len(cols)-1):.3f}'+r'\tabcolsep\relax}' for v in fractions)+'@{}'
        header=' & '.join(r'\textbf{'+esc(c)+'}' for c in cols)+r'\\\midrule'
        out.extend([r'\begingroup\small\setstretch{1.15}\setlength{\tabcolsep}{4pt}\renewcommand{\arraystretch}{1.2}',r'\renewcommand{\thetable}{'+number+r'}\setcounter{table}{0}',r'\begin{longtable}{'+spec+'}',r'\caption{'+esc(title)+r'}\label{tab:'+number+r'}\\',r'\toprule',header,r'\endfirsthead',r'\multicolumn{'+str(len(cols))+r'}{l}{\small Table '+number+r' continued}\\\toprule',header,r'\endhead',r'\bottomrule\endfoot'])
        for row in rows:
            assert len(row)==len(cols)
            out.append(' & '.join(esc(c) for c in row)+r'\\')
        out.append(r'\end{longtable}\endgroup'); tables+=1; i+=1; continue
    if line.startswith('[[FIGURE:'):
        number,title,kind=line[9:-2].split('|')
        out.extend([r'\begin{figure}[htbp]\centering',r'\resizebox{\linewidth}{!}{\begin{tikzpicture}',tikz(kind),r'\end{tikzpicture}}',r'\renewcommand{\thefigure}{'+number+r'}\setcounter{figure}{0}',r'\caption{'+esc(title)+r'}\label{fig:'+number+'}',r'\end{figure}'])
        figures+=1; i+=1; continue
    if line=='[[CODE]]':
        out.append(r'\begin{Verbatim}[fontsize=\small]'); i+=1
        while lines[i]!='[[ENDCODE]]': out.append(lines[i]); i+=1
        out.append(r'\end{Verbatim}'); i+=1; continue
    out.append(esc(line)); i+=1
out.append(r'\end{document}')
assert figures==7 and tables==31 and refs_added and len(refs)==33
OUT.write_text('\n'.join(out)+'\n',encoding='utf-8')
print(json.dumps({'output':str(OUT),'figures':figures,'tables':tables,'references':len(refs),'bytes':OUT.stat().st_size}))
