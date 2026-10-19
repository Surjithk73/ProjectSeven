"""Build the maintained academic DOCX from the retained university template.

Run with the bundled workspace Python. No application tests or model calls.
"""
from pathlib import Path
import copy
import hashlib
import io
import json
import re
import zipfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = Path(__file__).resolve().parent
QA = ROOT / 'reports' / '.qa'
QA.mkdir(exist_ok=True)
REF = Path(r'C:\Users\surji\Downloads\BCE497J PROJECT REPORT TEMPLATE_V1.docx')
OUT = ROOT / 'reports' / 'APIShield_Project_Report.docx'
TITLE = 'APIShield: Intent-Aware Detection and Validated Repair of Breaking API Changes'
BODY = (SRC / 'report.md').read_text(encoding='utf-8')
REFS = json.loads((SRC / 'references.json').read_text(encoding='utf-8'))

def box(ax, x, y, w, h, text, shade='#f2f2f2', size=11):
    patch = FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.015,rounding_size=0.025',
                          facecolor=shade,edgecolor='#333333',linewidth=0.9)
    ax.add_patch(patch)
    ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=size,fontfamily='DejaVu Sans')

def arrow(ax,a,b,dashed=False):
    ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='#333333',lw=1.1,
                linestyle='--' if dashed else '-'))

def diagram(kind):
    path=QA/f'figure_{kind}.png'
    fig,ax=plt.subplots(figsize=(7.4,2.8),dpi=260)
    ax.set_xlim(0,10);ax.set_ylim(0,4);ax.axis('off')
    if kind=='workflow':
        labels=['Connect project\nand contracts','Check changes\nand impact','Confirm rules\nand checks','Request repair\nand review','Patch or draft\npreview']
        for i,label in enumerate(labels):
            box(ax,.1+i*2,.95,1.75,1.6,label,size=10)
            if i<4:arrow(ax,(1.85+i*2,1.75),(2.08+i*2,1.75))
        ax.text(5,.25,'Explicit generation and publication actions remain separate',ha='center',fontsize=10)
    elif kind=='architecture':
        box(ax,.2,2.8,2.2,.8,'React browser');box(ax,.2,.6,2.2,.8,'CLI / research\nharness')
        box(ax,3.2,2.8,2.4,.8,'FastAPI coordination');box(ax,3.2,.6,2.4,.8,'Shared Python\nengine')
        box(ax,6.6,2.8,2.9,.8,'Durable local evidence');box(ax,6.6,.6,2.9,.8,'oasdiff / Gemini\nGitHub')
        arrow(ax,(2.4,3.2),(3.2,3.2));arrow(ax,(4.4,2.8),(4.4,1.4));arrow(ax,(2.4,1),(3.2,1))
        arrow(ax,(5.6,3.2),(6.6,3.2));arrow(ax,(5.6,1),(6.6,1));arrow(ax,(7.9,1.4),(7.9,2.8))
    elif kind=='states':
        box(ax,.2,1.5,1.5,.9,'Queued');box(ax,2.4,1.5,1.7,.9,'Running')
        outcomes=[('Completed',3.1),('Validated',2.05),('Needs input',1),('Failed / cancelled',-.05)]
        arrow(ax,(1.7,1.95),(2.4,1.95))
        for label,y in outcomes:
            box(ax,6.4,y+.2,2.8,.65,label,size=11);arrow(ax,(4.1,1.95),(6.4,y+.52))
        ax.text(3.6,.25,'No change is a separate repair outcome',fontsize=10,ha='center')
    elif kind=='graph':
        items=[(.2,2.8,2.1,.8,'Caller function'),(3.2,2.8,2.2,.8,'HTTP consumer'),(7,2.8,2.6,.8,'API operation'),
               (3.2,.5,2.2,.8,'Confirmed rule'),(7,.5,2.6,.8,'Committed test')]
        for x,y,w,h,t in items:box(ax,x,y,w,h,t)
        arrow(ax,(2.3,3.2),(3.2,3.2));arrow(ax,(5.4,3.2),(7,3.2));arrow(ax,(4.3,2.8),(4.3,1.3),True)
        arrow(ax,(5.4,.9),(7,.9),True)
        ax.text(2.7,3.65,'calls',fontsize=9,ha='center');ax.text(6.2,3.65,'consumes',fontsize=9,ha='center')
        ax.text(5,2.05,'declared scope',fontsize=9);ax.text(6.2,1.4,'checked by',fontsize=9,ha='center')
        ax.text(5,.05,'Solid arrows: static observation   Dashed arrows: developer declaration',ha='center',fontsize=9)
    elif kind=='publication':
        labels=['Validated\nretained patch','Offline\nrevalidation','Revisioned\nfull preview','Explicit\ndraft publish']
        for i,t in enumerate(labels):
            box(ax,.2+i*2.5,1.65,2,.95,t)
            if i<3:arrow(ax,(2.2+i*2.5,2.12),(2.65+i*2.5,2.12))
        box(ax,5.2,.1,4.5,.7,'Stale state or lost response → reconcile',size=10)
        arrow(ax,(8.7,1.65),(8.7,.8),True)
        ax.text(4.7,3.3,'Local acceptance and remote publication have distinct outcomes',ha='center',fontsize=11)
    elif kind=='pipeline':
        labels=['Pinned source\nand contracts','Static scan\nand comparison','Guarded\ncandidate edits','Configured tests\nand rule checks']
        for i,t in enumerate(labels):
            box(ax,.2+i*2.5,1.65,2,1,t)
            if i<3:arrow(ax,(2.2+i*2.5,2.15),(2.65+i*2.5,2.15))
        ax.text(5,.7,'Retry begins again from the same captured baseline',ha='center',fontsize=11)
        arrow(ax,(8.7,1.65),(1.2,1.65),True)
        ax.text(5,3.35,'Protected source, tests and contract integrity are checked',ha='center',fontsize=11)
    elif kind=='oracle':
        box(ax,.3,2.3,3.8,1,'Contracts and source\nRules and public feedback',size=11)
        box(ax,5.7,2.3,3.8,1,'Candidate generation\nand engine acceptance',size=11)
        arrow(ax,(4.1,2.8),(5.7,2.8));box(ax,5.7,.35,3.8,1,'Separate verification clone\nHidden behavioral checks',size=11)
        arrow(ax,(7.6,2.3),(7.6,1.35));ax.plot([.4,9.6],[1.75,1.75],ls='--',color='#777777')
        ax.text(2.5,.7,'Hidden checks never enter\nprompts or retry feedback',ha='center',fontsize=11)
    for text_item in ax.texts:
        text_item.set_fontsize(max(11.5,text_item.get_fontsize()+1.5))
    fig.tight_layout(pad=.15);fig.savefig(path,bbox_inches='tight',facecolor='white');plt.close(fig)
    return path

doc=Document(REF)
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(3.81);sec.right_margin=Cm(2.54);sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54)
sec.header_distance=Cm(1.27);sec.footer_distance=Cm(1.27)
sec.different_first_page_header_footer=False
for area in (sec.header,sec.footer,sec.first_page_header,sec.first_page_footer):
    for node in list(area._element):area._element.remove(node)
for tag in ('w:pgNumType','w:titlePg'):
    for node in sec._sectPr.findall(qn(tag)):sec._sectPr.remove(node)
num=OxmlElement('w:pgNumType');num.set(qn('w:fmt'),'decimal');num.set(qn('w:start'),'1');sec._sectPr.append(num)

def style(name,size=12,bold=False,space=6,line=1.5):
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    rpr=s.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Times New Roman')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+k),None)
    pf=s.paragraph_format;pf.space_after=Pt(space);pf.space_before=Pt(0);pf.line_spacing=line;pf.widow_control=True
    pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(0)
    return s

style('Normal',space=4);style('Title',20,True,14,1.25)
style('Heading 1',16,True,16,1.25);style('Heading 2',14,False,10,1.25);style('Heading 3',12,False,8,1.25)
style('Caption',10.5,False,6,1);style('Report Table',12,False,0,1.5);style('Report Code',12,False,0,1.5)
style('Report TOC',12,False,2,1);style('Report Reference',12,False,12,1.5)
style('Front Heading',14,True,18,1.25)
for name in ['Heading 1','Heading 2','Heading 3','Front Heading']:
    doc.styles[name].paragraph_format.keep_with_next=True
for s in doc.styles:
    if s.type==WD_STYLE_TYPE.PARAGRAPH:
        try:s.font.color.rgb=RGBColor(0,0,0)
        except ValueError:pass

settings=doc.settings.element
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)

def field(p,instruction,placeholder='1'):
    run=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');run._r.append(b)
    run=p.add_run();t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';run._r.append(t)
    run=p.add_run();s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');run._r.append(s)
    p.add_run(placeholder)
    run=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');run._r.append(e)

foot=sec.footer.add_paragraph();foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.space_after=Pt(0);field(foot,'PAGE')
for r in foot.runs:r.font.name='Times New Roman';r.font.size=Pt(12)

bookmark_id=0
def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmark_id));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bookmark_id))
    p._p.insert(0,a);p._p.append(b)

pending_page=False
def paragraph(text='',sty=None,align=None,size=None,bold=None,italic=None):
    global pending_page
    p=doc.add_paragraph(style=sty)
    if pending_page:
        p.paragraph_format.page_break_before=True
        pending_page=False
    r=p.add_run(text)
    if align is not None:p.alignment=align
    elif sty is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    elif sty in ('Heading 2','Heading 3'):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if size:r.font.size=Pt(size)
    if bold is not None:r.bold=bold
    if italic is not None:r.italic=italic
    return p

def newpage():
    global pending_page
    pending_page=True

def heading(text,front=False,mark=None):
    p=paragraph(text.upper(),'Front Heading' if front else 'Heading 1',WD_ALIGN_PARAGRAPH.CENTER)
    if front:
        for r in p.runs:r.underline=True
    if mark:bookmark(p,mark)
    return p

with zipfile.ZipFile(REF) as z:logo=z.read('word/media/image2.jpg')
logo_path=QA/'institution_logo.jpg';logo_path.write_bytes(logo)

def titlepage(cover=False):
    p=paragraph('A project report on' if cover else 'PROJECT REPORT',align=WD_ALIGN_PARAGRAPH.CENTER,size=12,italic=True)
    p.paragraph_format.space_before=Pt(20)
    paragraph(TITLE.upper(),'Title',WD_ALIGN_PARAGRAPH.CENTER)
    paragraph('Submitted in partial fulfillment for the award of the degree of',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,italic=True)
    paragraph('Bachelor of Technology\nin\n[PROGRAM NAME AND SPECIALIZATION]',align=WD_ALIGN_PARAGRAPH.CENTER,size=22,bold=True).paragraph_format.line_spacing=1.15
    paragraph('by',align=WD_ALIGN_PARAGRAPH.CENTER,size=14,italic=True)
    p=paragraph('[CANDIDATE NAME] ([REGISTRATION NUMBER])\n[TEAM MEMBER 2 NAME AND REGISTRATION NUMBER]\n[TEAM MEMBER 3 NAME AND REGISTRATION NUMBER]',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,bold=True)
    p.paragraph_format.line_spacing=1.15
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(logo_path),width=Inches(2.65))
    paragraph('SCHOOL OF COMPUTER SCIENCE AND ENGINEERING',align=WD_ALIGN_PARAGRAPH.CENTER,size=16,bold=True)
    paragraph('VELLORE INSTITUTE OF TECHNOLOGY\nCHENNAI',align=WD_ALIGN_PARAGRAPH.CENTER,size=12,bold=True)
    paragraph('[MONTH AND YEAR OF SUBMISSION]',align=WD_ALIGN_PARAGRAPH.CENTER,size=12)

titlepage(True);newpage();titlepage();newpage()
heading('Declaration',True,'front_declaration')
paragraph('I hereby declare that the project report entitled “'+TITLE+'”, submitted by [CANDIDATE NAME] ([REGISTRATION NUMBER]) for the award of the degree of Bachelor of Technology in [PROGRAM NAME AND SPECIALIZATION] at Vellore Institute of Technology, Chennai, is a record of bonafide work carried out by me under the supervision of [GUIDE NAME AND DESIGNATION].')
paragraph('I further declare that the work reported in this thesis has not been submitted and will not be submitted, either in part or in full, for the award of any other degree or diploma in this institute or any other institute or university.')
paragraph('Place: Chennai\nDate: [DATE]',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True).paragraph_format.space_before=Pt(48)
paragraph('Signature of the Candidate: ____________________\n[CANDIDATE NAME]\n[REGISTRATION NUMBER]',align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)
newpage();heading('Certificate',True,'front_certificate')
paragraph('School of Computer Science and Engineering',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
paragraph('This is to certify that the report entitled “'+TITLE+'” is prepared and submitted by [CANDIDATE NAME] ([REGISTRATION NUMBER]) to Vellore Institute of Technology, Chennai, in partial fulfillment of the requirements for the award of the degree of Bachelor of Technology in [PROGRAM NAME AND SPECIALIZATION], and is a bonafide record of work carried out under my guidance. The project fulfills the requirements as per the regulations of this University and, in my opinion, meets the necessary standards for submission. The contents of this report have not been submitted and will not be submitted, either in part or in full, for the award of any other degree or diploma, and the same is certified.')
paragraph('Signature of the Guide: ____________________\nName: [GUIDE NAME AND DESIGNATION]\nDate: [DATE]',align=WD_ALIGN_PARAGRAPH.LEFT)
paragraph('Signature of Examiner 1: ____________________\nName: [EXAMINER NAME]    Date: [DATE]\n\nSignature of Examiner 2: ____________________\nName: [EXAMINER NAME]    Date: [DATE]',align=WD_ALIGN_PARAGRAPH.LEFT)
paragraph('Approved by the Head of Department\n[PROGRAM NAME AND SPECIALIZATION]\nName: [HOD NAME]    Date: [DATE]\n[SEAL OF THE SCHOOL]',align=WD_ALIGN_PARAGRAPH.CENTER)
newpage();heading('Abstract',True,'front_abstract')
paragraph('Changes to a provider API can leave a client structurally incompatible or alter the business behavior of an apparently successful migration. APIShield is a local application for detecting documented OpenAPI changes, identifying supported consuming Python calls and preparing tested migration proposals. It connects a pinned source revision and contract snapshots to developer-confirmed requirements and executable checks. Candidates are generated or supplied in a disposable workspace, constrained to affected source files and accepted only after configured validation and applicable requirement checks succeed.')
paragraph('The system combines a shared Python engine, FastAPI coordination and a React workspace. Captured source and paired contracts supply a structural graph; inferred intent produces additional executable preservation checks. Reviewed repairs can be applied with freshness checks and recovery records. Historical acceptance demonstrates selected endpoint, payload, response, refund and exact-integer repairs alongside provider failures and a reviewed missing-information diagnostic.')
paragraph('The retained research package implements six benchmark configurations. Twenty-four authored-candidate qualification executions across four tasks distinguish twelve public-test-passing but behaviorally incorrect candidates from four passing candidates, four syntax failures and four abstentions. Eighteen live inference traces across three passes record twelve successful provider responses; their model-generated candidates were not evaluated by public or hidden tests. The evidence establishes evaluator discrimination and observed inference costs, while comparative repair effectiveness on a representative 120-task corpus remains unmeasured. Limitations include narrow fixtures, partial static coverage, provider variability and incomplete full-study execution.')
paragraph('Keywords: API evolution; OpenAPI; program repair; behavioral requirements; repository graph; validation provenance.',size=12,italic=True)
newpage();heading('Acknowledgement',True,'front_acknowledgement')
paragraph('I express my sincere gratitude to [GUIDE NAME], [DESIGNATION], School of Computer Science and Engineering, Vellore Institute of Technology, Chennai, for the guidance and feedback provided during this project. The discussions helped me refine the problem, distinguish implementation evidence from research claims and identify the work required for a defensible evaluation.')
paragraph('I thank [CHANCELLOR NAME], [VICE-CHANCELLOR NAME], [PRO-VICE-CHANCELLOR NAME] and the university administration for the academic environment and facilities made available during the programme. I also thank [DEAN NAME], [HOD NAME] and [PROJECT COORDINATOR NAMES] for their support and coordination.')
paragraph('I am grateful to the faculty and staff of the School of Computer Science and Engineering for their teaching and assistance. I acknowledge the authors and maintainers of the research and open-source tools cited in this report. Their work provides both the technical foundation and the comparisons needed to assess this project fairly.')
paragraph('I thank my family, teammates and friends for their encouragement throughout the work.')
paragraph('Place: Chennai\nDate: [DATE]\n\n[CANDIDATE NAME]',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True)

# Inventory headings and figure/table captions before writing their navigators.
pages=BODY.split('<!-- PAGE -->')
entries=[];figures=[];tables=[]
for pi,chunk in enumerate(pages):
    lines=chunk.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('# CHAPTER ') or line.startswith('# APPENDIX '):
            label=line[2:];title=lines[i+1][2:] if i+1<len(lines) else ''
            entries.append((label+' '+title,f'section_{pi}_{i}',0));i+=2;continue
        if line.startswith('# REFERENCES'):entries.append(('REFERENCES',f'section_{pi}_{i}',0))
        elif line.startswith('## '):entries.append((line[3:].upper(),f'section_{pi}_{i}',1))
        elif line.startswith('[[FIGURE:'):
            number,caption,kind=line[len('[[FIGURE:'):-2].split('|');figures.append((number,caption,'fig_'+number.replace('.','_')))
        elif line.startswith('[[TABLE:'):
            vals=line[len('[[TABLE:'):-2].split('|');tables.append((vals[0],vals[1],'tbl_'+vals[0].replace('.','_')))
        i+=1

def navline(label,mark,indent=False):
    if len(label)>68:
        at=label.rfind(' ',0,65)
        if at>0:label=label[:at]+'\n'+label[at+1:]
    p=paragraph(label,'Report TOC');p.paragraph_format.left_indent=Cm(.4 if indent else 0)
    p.paragraph_format.keep_together=True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    p.add_run('\t');field(p,'PAGEREF '+mark+' \\h')
    return p

front_entries=[('DECLARATION','front_declaration',0),('CERTIFICATE','front_certificate',0),('ABSTRACT','front_abstract',0),('ACKNOWLEDGEMENT','front_acknowledgement',0),('LIST OF FIGURES','front_figures',0),('LIST OF TABLES','front_tables',0),('LIST OF SYMBOLS AND ABBREVIATIONS','front_abbreviations',0)]
toc=front_entries+entries
split=(len(toc)+1)//2
newpage();heading('Table of contents',True)
for label,mark,level in toc[:split]:navline(label,mark,level)
newpage();heading('Table of contents continued',True)
for label,mark,level in toc[split:]:navline(label,mark,level)
newpage();heading('List of figures',True,'front_figures')
for number,caption,mark in figures:navline(number+'  '+caption,mark)
newpage();heading('List of tables',True,'front_tables')
for number,caption,mark in tables:navline(number+'  '+caption,mark)
newpage();heading('List of symbols and abbreviations',True,'front_abbreviations')
abbreviations=[('API','Application programming interface'),('AST','Abstract syntax tree'),('BFF','Backend for frontend'),('CI','Continuous integration'),('CLI','Command-line interface'),('HTTP','Hypertext Transfer Protocol'),('IPS','Intent Preservation Score; optional legacy model judgment'),('JSON','JavaScript Object Notation'),('LLM','Large language model'),('MAPE-K','Monitor, analyze, plan, execute over knowledge'),('OAS','OpenAPI Specification'),('PR','Pull request'),('RQ','Research question'),('SHA','Secure Hash Algorithm'),('SIG','Semantic Intent Graph'),('UI','User interface'),('B1–B6','Proposed comparative research configurations')]
for k,v in abbreviations:
    p=paragraph(k+'\t'+v);p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.2))

def caption(text,mark):
    p=paragraph(text,'Caption',WD_ALIGN_PARAGRAPH.CENTER);bookmark(p,mark)
    p.paragraph_format.keep_with_next=True
    return p

def table(number,title,columns,rows):
    caption('Table '+number+'  '+title,'tbl_'+number.replace('.','_'))
    t=doc.add_table(rows=1,cols=len(columns));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    total=14.65
    fractions={2:[.34,.66],3:[.27,.38,.35],4:[.27,.24,.27,.22]}[len(columns)]
    if number=='A6.1':fractions=[.52,.48]
    for c,f in zip(t.columns,fractions):c.width=Cm(total*f)
    for ci,col in enumerate(columns):t.rows[0].cells[ci].text=col
    for row in rows:
        cells=t.add_row().cells
        for ci,text in enumerate(row):cells[ci].text=text
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(total*fractions[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'B8B8B8');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'35' if edge in ['top','bottom'] else '75');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');tcpr.append(sh)
            for p in cell.paragraphs:
                p.style=doc.styles['Report Table'];p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                for r in p.runs:r.bold=ri==0
    spacer=paragraph('');spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.line_spacing=Pt(2)
    for r in spacer.runs:r.font.size=Pt(1)

for pi,chunk in enumerate(pages):
    lines=chunk.strip().splitlines()
    if pi==0 or lines[0].startswith('# ') or lines[0].startswith('[[REFERENCES:'):
        newpage()
    else:
        # Let sections flow naturally; only chapters and appendices require new pages.
        doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    i=0;buf=[]
    def flush():
        if buf:paragraph(' '.join(buf));buf.clear()
    while i<len(lines):
        line=lines[i].strip()
        if not line:flush();i+=1;continue
        if line.startswith('# CHAPTER ') or line.startswith('# APPENDIX '):
            flush();text=line[2:]+'\n'+lines[i+1][2:];heading(text,mark=f'section_{pi}_{i}');i+=2;continue
        if line.startswith('# '):
            flush();heading(line[2:],mark=f'section_{pi}_{i}');i+=1;continue
        if line.startswith('## '):
            flush();p=paragraph(line[3:].upper(),'Heading 2');bookmark(p,f'section_{pi}_{i}');i+=1;continue
        if line.startswith('[[TABLE:'):
            flush();vals=line[len('[[TABLE:'):-2].split('|');rows=[];i+=1
            while not lines[i].startswith('[[ENDTABLE]]'):rows.append(lines[i].split('|'));i+=1
            table(vals[0],vals[1],vals[2:],rows);i+=1;continue
        if line.startswith('[[FIGURE:'):
            flush();number,title,kind=line[len('[[FIGURE:'):-2].split('|')
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(diagram(kind)),width=Cm(14.5));caption('Figure '+number+'  '+title,'fig_'+number.replace('.','_')).paragraph_format.keep_with_next=False
            i+=1;continue
        if line.startswith('[[CODE]]'):
            flush();i+=1
            while not lines[i].startswith('[[ENDCODE]]'):
                p=paragraph(lines[i],'Report Code');p.paragraph_format.keep_with_next=False;i+=1
            paragraph('');i+=1;continue
        if line.startswith('[[REFERENCES:'):
            flush();batch=int(re.search(r':(\d+)',line).group(1));ranges=[(0,6),(6,12),(12,18),(18,len(REFS))]
            a,b=ranges[batch-1]
            if batch>1:paragraph('REFERENCES CONTINUED','Heading 2')
            for ref in REFS[a:b]:
                p=paragraph(sty='Report Reference');p.paragraph_format.left_indent=Cm(1.27);p.paragraph_format.first_line_indent=Cm(-1.27)
                p.add_run(ref['author'].rstrip('.')+'. ('+ref['year']+'). ')
                article=ref['author'].startswith(('Atlidakis','Debenedetti','Kephart','Zhan,'))
                p.add_run(ref['title']+'. ').italic=not article
                publication=ref['type']
                if publication.startswith(('Unpublished','Software documentation','Source repository','Technical specification')):
                    p.add_run('['+publication+']. ')
                else:p.add_run(publication+'. ').italic=article
                if ref['url'].startswith('http'):p.add_run(ref['url'])
                else:p.add_run('ProjectSeven repository, '+ref['url']+'.')
            i+=1;continue
        buf.append(line);i+=1
    flush()

# Retain a task-local audit rather than changing project application state.
doc.core_properties.title=TITLE;doc.core_properties.subject='BCE497J final project report'
doc.core_properties.author='[CANDIDATE NAME]';doc.core_properties.keywords='APIShield, OpenAPI, program repair, requirements'
doc.core_properties.comments=''
doc.save(OUT)
manifest={'title':TITLE,'output':str(OUT),'reference':str(REF),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),
          'planned_pages':len(pages)+11,'source_word_count':len(BODY.split()),'figures':len(figures),'tables':len(tables),'references':len(REFS),
          'evidence_cutoff':'2026-10-15','tests_run_for_report':False,'controlled_results':'Four-task authored qualification and live inference observations; full model correctness benchmark pending',
          'placeholders':'Personal details, real signatures, submission date, controlled-study measurements',
          'rendering':'Native Word field refresh and PDF export; PyMuPDF raster QA',
          'sources':[str(SRC/'report.md'),str(SRC/'references.json')]}
(SRC/'report_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))
